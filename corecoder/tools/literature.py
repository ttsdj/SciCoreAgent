"""Literature search and red-blue review tools."""

from __future__ import annotations

import json
import re
from typing import Any

from .base import Tool


def _json(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2)


def _load_jsonish(value: Any) -> Any:
    parsed, _ = _parse_json_with_fallbacks(value)
    return parsed


def _parse_json_with_fallbacks(value: Any) -> tuple[Any, list[dict[str, str]]]:
    """Parse literature JSON using a three-layer fallback chain."""
    fallbacks = []
    if isinstance(value, (dict, list)):
        return value, [{"layer": "direct_object", "status": "success"}]
    if value in (None, ""):
        return [], [{"layer": "empty_input", "status": "success"}]

    text = str(value).strip()
    try:
        return json.loads(text), [{"layer": "direct_json", "status": "success"}]
    except Exception as exc:
        fallbacks.append({"layer": "direct_json", "status": "failed", "reason": str(exc)})

    fenced = re.search(r"```(?:json)?\s*(.*?)```", text, re.S | re.I)
    if fenced:
        try:
            return json.loads(fenced.group(1).strip()), [*fallbacks, {"layer": "fenced_json", "status": "success"}]
        except Exception as exc:
            fallbacks.append({"layer": "fenced_json", "status": "failed", "reason": str(exc)})

    start_positions = [pos for pos in (text.find("["), text.find("{")) if pos >= 0]
    end_positions = [pos for pos in (text.rfind("]"), text.rfind("}")) if pos >= 0]
    if start_positions and end_positions:
        snippet = text[min(start_positions): max(end_positions) + 1]
        try:
            return json.loads(snippet), [*fallbacks, {"layer": "substring_json", "status": "success"}]
        except Exception as exc:
            fallbacks.append({"layer": "substring_json", "status": "failed", "reason": str(exc)})

    rows = []
    for line in text.splitlines():
        line = line.strip().rstrip(",")
        if not line:
            continue
        try:
            item = json.loads(line)
        except Exception:
            continue
        if isinstance(item, dict):
            rows.append(item)
    if rows:
        return rows, [*fallbacks, {"layer": "json_lines", "status": "success"}]

    fallbacks.append({"layer": "json_lines", "status": "failed", "reason": "no parseable JSON lines"})
    raise ValueError("items_json could not be parsed by any JSON fallback")


class PubMedSearchTool(Tool):
    name = "pubmed_search"
    description = "Search PubMed via the bundled PubMed literature MCP implementation. Requires NCBI_EMAIL."
    parameters = {
        "type": "object",
        "properties": {
            "query": {"type": "string", "description": "PubMed search query."},
            "max_results": {"type": "integer", "description": "Maximum result count.", "default": 20},
            "date_from": {"type": "string", "description": "Optional start date YYYY-MM-DD."},
            "date_to": {"type": "string", "description": "Optional end date YYYY-MM-DD."},
            "sort": {"type": "string", "description": "relevance or pub_date.", "default": "relevance"},
        },
        "required": ["query"],
    }

    def execute(
        self,
        query: str,
        max_results: int = 20,
        date_from: str = "",
        date_to: str = "",
        sort: str = "relevance",
    ) -> str:
        from mcp_servers.pubmed_literature.server import PubMedServer

        server = PubMedServer()
        result = server.pubmed_search(
            query=query,
            max_results=int(max_results or 20),
            date_from=date_from or None,
            date_to=date_to or None,
            sort=sort or "relevance",
        )
        result["status"] = server.status()
        return _json(result)


class PubMedFetchDetailsTool(Tool):
    name = "pubmed_fetch_details"
    description = "Fetch PubMed article metadata and abstracts for PMIDs. Requires NCBI_EMAIL."
    parameters = {
        "type": "object",
        "properties": {
            "pmids": {
                "type": "array",
                "items": {"type": "string"},
                "description": "PMIDs to fetch.",
            },
            "include_abstract": {"type": "boolean", "default": True},
        },
        "required": ["pmids"],
    }

    def execute(self, pmids: list[str], include_abstract: bool = True) -> str:
        from mcp_servers.pubmed_literature.server import PubMedServer

        server = PubMedServer()
        return _json(server.pubmed_fetch_details(pmids=pmids, include_abstract=include_abstract))


class PubMedLiteratureReviewTool(Tool):
    name = "pubmed_literature_review"
    description = "Run the bundled one-stop PubMed literature review workflow and return generated report paths."
    parameters = {
        "type": "object",
        "properties": {
            "query": {"type": "string", "description": "PubMed query."},
            "topic": {"type": "string", "description": "Review topic label.", "default": "literature review"},
            "max_results": {"type": "integer", "default": 20},
            "output_docx": {"type": "string", "description": "Output docx path.", "default": "reports/literature_review.docx"},
        },
        "required": ["query"],
    }

    def execute(
        self,
        query: str,
        topic: str = "literature review",
        max_results: int = 20,
        output_docx: str = "reports/literature_review.docx",
    ) -> str:
        from mcp_servers.pubmed_literature.server import PubMedServer

        server = PubMedServer()
        return _json(
            server.pubmed_literature_review(
                topic=topic or "literature review",
                query=query,
                max_results=int(max_results or 20),
                output_docx=output_docx or "reports/literature_review.docx",
            )
        )


class LiteratureRedBlueReviewTool(Tool):
    name = "literature_red_blue_review"
    description = (
        "Run a full Red-Blue-Judge evidence audit over literature rows. "
        "Red Agent attacks factuality, logic, citation quality, and bias; Blue Agent applies "
        "ADD/DELETE/MODIFY/VERIFY repairs; Judge scores convergence and detects oscillation."
    )
    parameters = {
        "type": "object",
        "properties": {
            "items_json": {
                "type": "string",
                "description": "JSON list of literature rows/articles to audit.",
            },
            "claim": {"type": "string", "description": "Main review claim or question.", "default": ""},
            "iterations": {"type": "integer", "description": "Maximum audit iterations.", "default": 2},
            "convergence_threshold": {
                "type": "integer",
                "description": "Stop when judge score reaches this value.",
                "default": 90,
            },
        },
        "required": ["items_json"],
    }

    def execute(
        self,
        items_json: str,
        claim: str = "",
        iterations: int = 2,
        convergence_threshold: int = 90,
    ) -> str:
        try:
            parsed, parser_fallbacks = _parse_json_with_fallbacks(items_json)
        except Exception as exc:
            return _json({"error": str(exc), "parser_fallbacks": [{"layer": "all", "status": "failed"}]})

        items = parsed
        if isinstance(items, dict):
            items = items.get("articles") or items.get("rows") or [items]
        if not isinstance(items, list):
            return _json({"error": "items_json must decode to a list or object."})

        state = RedBlueState(items=[_normalize_item(item, index) for index, item in enumerate(items)])
        red_agent = RedAgent()
        blue_agent = BlueAgent()
        judge_agent = JudgeAgent(threshold=int(convergence_threshold or 90))
        max_iter = max(1, min(int(iterations or 2), 8))
        signatures = []
        round_reports = []
        oscillation = False
        converged = False
        stop_reason = "max_iterations"

        for round_index in range(1, max_iter + 1):
            attacks = red_agent.attack(state.items, claim=claim)
            repairs = blue_agent.repair(state.items, attacks)
            state.items = blue_agent.apply(state.items, repairs)
            judgment = judge_agent.score(state.items, attacks, repairs)
            signature = _state_signature(state.items, attacks)
            if signature in signatures:
                oscillation = True
            signatures.append(signature)
            round_reports.append(
                {
                    "round": round_index,
                    "red_agent_attacks": attacks,
                    "blue_agent_repairs": repairs,
                    "judge": judgment,
                }
            )
            if judgment["score"] >= judge_agent.threshold and not attacks:
                converged = True
                stop_reason = "score_converged"
                break
            if oscillation:
                stop_reason = "oscillation_detected"
                break

        final_score = round_reports[-1]["judge"]["score"] if round_reports else 0
        score_history = [round_item["judge"]["score"] for round_item in round_reports]
        return _json(
            {
                "claim": claim,
                "item_count": len(state.items),
                "score": final_score,
                "converged": converged,
                "oscillation_detected": oscillation,
                "stop_reason": stop_reason,
                "convergence_trace": {
                    "score_history": score_history,
                    "score_delta": (score_history[-1] - score_history[0]) if len(score_history) >= 2 else 0,
                    "threshold": judge_agent.threshold,
                },
                "rounds": round_reports,
                "final_items": state.items,
                "parser_fallbacks": parser_fallbacks,
            }
        )


class LiteratureExportXlsxTool(Tool):
    name = "literature_export_xlsx"
    description = "Export literature rows and optional Red-Blue review output to a stable .xlsx workbook."
    parameters = {
        "type": "object",
        "properties": {
            "items_json": {"type": "string", "description": "JSON list of literature rows."},
            "output_path": {"type": "string", "description": "Output .xlsx path inside workspace."},
            "red_blue_json": {"type": "string", "description": "Optional JSON output from literature_red_blue_review."},
        },
        "required": ["items_json", "output_path"],
    }

    def execute(self, items_json: str, output_path: str, red_blue_json: str = "") -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except Exception as exc:
            return f"Error: openpyxl is required for xlsx export: {exc}"

        items, item_fallbacks = _parse_json_with_fallbacks(items_json)
        if isinstance(items, dict):
            items = items.get("final_items") or items.get("articles") or items.get("rows") or [items]
        if not isinstance(items, list):
            return "Error: items_json must parse to a list or object"

        red_blue = {}
        red_blue_fallbacks = []
        if red_blue_json:
            try:
                red_blue, red_blue_fallbacks = _parse_json_with_fallbacks(red_blue_json)
            except Exception as exc:
                red_blue = {"parse_error": str(exc)}

        wb = Workbook()
        ws = wb.active
        ws.title = "Literature"
        normalized_items = [_normalize_item(item, index) for index, item in enumerate(items)]
        headers = [
            "title", "year", "journal", "doi", "pmid", "core_conclusion",
            "evidence_strength", "limitations", "verification_tasks", "red_blue_notes",
        ]
        _write_table(ws, headers, normalized_items)

        if isinstance(red_blue, dict) and red_blue:
            summary = wb.create_sheet("RedBlue Summary")
            summary_rows = [
                {"field": "score", "value": red_blue.get("score", "")},
                {"field": "converged", "value": red_blue.get("converged", "")},
                {"field": "oscillation_detected", "value": red_blue.get("oscillation_detected", "")},
                {"field": "stop_reason", "value": red_blue.get("stop_reason", "")},
                {"field": "score_history", "value": json.dumps(red_blue.get("convergence_trace", {}).get("score_history", []), ensure_ascii=False)},
            ]
            _write_table(summary, ["field", "value"], summary_rows)

            rounds = wb.create_sheet("RedBlue Rounds")
            round_rows = []
            for round_item in red_blue.get("rounds", []):
                round_no = round_item.get("round", "")
                for attack in round_item.get("red_agent_attacks", []):
                    round_rows.append({"round": round_no, "type": "red_attack", **attack})
                for repair in round_item.get("blue_agent_repairs", []):
                    round_rows.append({"round": round_no, "type": "blue_repair", **repair})
            _write_table(rounds, ["round", "type", "row", "dimension", "issue", "severity", "action", "field", "instruction"], round_rows)

        meta = wb.create_sheet("Metadata")
        _write_table(
            meta,
            ["field", "value"],
            [
                {"field": "item_parser_fallbacks", "value": json.dumps(item_fallbacks, ensure_ascii=False)},
                {"field": "red_blue_parser_fallbacks", "value": json.dumps(red_blue_fallbacks, ensure_ascii=False)},
            ],
        )

        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = min(max(len(str(col[0].value or "")) + 4, 16), 48)

        from pathlib import Path

        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        return _json({"output_path": str(target), "sheet_count": len(wb.worksheets), "item_count": len(normalized_items)})


class RedBlueState:
    def __init__(self, items: list[dict[str, Any]]):
        self.items = items


class RedAgent:
    def attack(self, items: list[dict[str, Any]], claim: str = "") -> list[dict[str, Any]]:
        attacks = []
        for index, item in enumerate(items):
            if item.get("_deleted"):
                continue
            title = item.get("title", "")
            year = item.get("year", "")
            doi = item.get("doi", "")
            pmid = item.get("pmid", "")
            abstract = item.get("abstract", "")
            conclusion = item.get("core_conclusion", "")
            evidence_strength = item.get("evidence_strength", "")
            limitations = item.get("limitations", "")

            if not title:
                attacks.append(_attack(index, "factuality", "Missing title.", "high", "ADD"))
            if not year:
                attacks.append(_attack(index, "factuality", "Missing publication year.", "medium", "ADD"))
            if doi and not _looks_like_doi(doi):
                attacks.append(_attack(index, "citation_quality", "DOI format looks invalid.", "medium", "VERIFY"))
            if not doi and not pmid:
                attacks.append(_attack(index, "citation_quality", "Missing DOI or PMID; verify citation identity.", "medium", "ADD"))
            if not abstract and not conclusion:
                attacks.append(_attack(index, "logic_consistency", "No abstract or conclusion text supports this row.", "high", "VERIFY"))
            if _overclaims(conclusion):
                attacks.append(_attack(index, "logic_consistency", "Conclusion uses over-strong causal language.", "medium", "MODIFY"))
            if conclusion and not evidence_strength:
                attacks.append(_attack(index, "evidence_strength", "Evidence strength is not classified.", "medium", "ADD"))
            if not limitations:
                attacks.append(_attack(index, "bias_and_limits", "Limitations are missing.", "medium", "ADD"))
            if claim and title and not _shares_terms(claim, title + " " + abstract + " " + conclusion):
                attacks.append(_attack(index, "relevance", "Row may be weakly related to the review claim.", "low", "VERIFY"))
        return attacks


class BlueAgent:
    def repair(self, items: list[dict[str, Any]], attacks: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return [_repair(attack, items[attack["row"]] if attack["row"] < len(items) else {}) for attack in attacks]

    def apply(self, items: list[dict[str, Any]], repairs: list[dict[str, Any]]) -> list[dict[str, Any]]:
        repaired = [dict(item) for item in items]
        for repair in repairs:
            row = repair["row"]
            if row >= len(repaired):
                continue
            item = repaired[row]
            action = repair["action"]
            field = repair.get("field", "")
            if action == "DELETE":
                item["_deleted"] = True
                item.setdefault("red_blue_notes", []).append(repair["instruction"])
            elif action == "MODIFY" and field:
                item[field] = _soften_claim(str(item.get(field, "")))
                item.setdefault("red_blue_notes", []).append(repair["instruction"])
            elif action == "ADD" and field:
                item[field] = item.get(field) or repair.get("placeholder", "VERIFY_REQUIRED")
                item.setdefault("red_blue_notes", []).append(repair["instruction"])
            else:
                item.setdefault("verification_tasks", []).append(repair["instruction"])
        return repaired


class JudgeAgent:
    def __init__(self, threshold: int = 90):
        self.threshold = threshold

    def score(
        self,
        items: list[dict[str, Any]],
        attacks: list[dict[str, Any]],
        repairs: list[dict[str, Any]],
    ) -> dict[str, Any]:
        penalty = 0
        for attack in attacks:
            penalty += {"high": 20, "medium": 10, "low": 5}.get(attack["severity"], 10)
        unresolved = [
            attack for attack in attacks
            if attack["preferred_action"] in {"VERIFY", "ADD"} and attack["severity"] in {"high", "medium"}
        ]
        score = max(0, min(100, 100 - penalty))
        return {
            "score": score,
            "threshold": self.threshold,
            "passed": score >= self.threshold and not unresolved,
            "unresolved_count": len(unresolved),
            "repair_count": len(repairs),
            "active_item_count": len([item for item in items if not item.get("_deleted")]),
        }


def _normalize_item(item: Any, index: int) -> dict[str, Any]:
    if not isinstance(item, dict):
        return {"_row": index, "_raw": item, "_deleted": True, "red_blue_notes": ["DELETE non-object row."]}
    aliases = {
        "Title": "title",
        "Year": "year",
        "DOI": "doi",
        "PMID": "pmid",
        "Abstract": "abstract",
        "Core Conclusion": "core_conclusion",
        "Evidence Strength": "evidence_strength",
        "Limitations": "limitations",
    }
    normalized = {"_row": index}
    for key, value in item.items():
        normalized[aliases.get(key, key)] = value
    for key in ("title", "year", "doi", "pmid", "abstract", "core_conclusion", "evidence_strength", "limitations"):
        normalized[key] = str(normalized.get(key, "") or "").strip()
    return normalized


def _attack(row: int, dimension: str, issue: str, severity: str, preferred_action: str) -> dict:
    return {
        "row": row,
        "dimension": dimension,
        "issue": issue,
        "severity": severity,
        "preferred_action": preferred_action,
    }


def _repair(attack: dict, item: dict[str, Any]) -> dict:
    issue = attack["issue"].lower()
    action = attack.get("preferred_action", "VERIFY")
    field = ""
    placeholder = "VERIFY_REQUIRED"
    if "title" in issue:
        field = "title"
        placeholder = "VERIFY_TITLE"
    elif "year" in issue:
        field = "year"
        placeholder = "VERIFY_YEAR"
    elif "doi" in issue:
        field = "doi"
        placeholder = "VERIFY_DOI_OR_PMID"
    elif "evidence strength" in issue:
        field = "evidence_strength"
        placeholder = "VERIFY_AND_CLASSIFY"
    elif "limitations" in issue:
        field = "limitations"
        placeholder = "VERIFY_LIMITATIONS"
    elif "over-strong" in issue:
        field = "core_conclusion"
    return {
        "row": attack["row"],
        "action": action,
        "target_dimension": attack["dimension"],
        "field": field,
        "placeholder": placeholder,
        "instruction": f"{action} evidence or wording to fix: {attack['issue']}",
    }


def _looks_like_doi(value: str) -> bool:
    return bool(re.match(r"^10\.\d{4,9}/\S+$", value.strip()))


def _overclaims(text: str) -> bool:
    return any(term in text.lower() for term in ["proves", "definitive", "always", "never", "完全证明", "必然"])


def _soften_claim(text: str) -> str:
    replacements = {
        "proves": "supports",
        "Proves": "Supports",
        "definitive": "strong",
        "always": "often",
        "never": "rarely",
        "完全证明": "支持",
        "必然": "可能",
    }
    result = text
    for old, new in replacements.items():
        result = result.replace(old, new)
    return result or "VERIFY_CLAIM_WORDING"


def _shares_terms(claim: str, evidence: str) -> bool:
    terms = {term for term in re.findall(r"[A-Za-z0-9]+", claim.lower()) if len(term) >= 4}
    if not terms:
        return True
    evidence_lower = evidence.lower()
    return any(term in evidence_lower for term in terms)


def _state_signature(items: list[dict[str, Any]], attacks: list[dict[str, Any]]) -> str:
    payload = {
        "items": [
            {
                "row": item.get("_row"),
                "doi": item.get("doi"),
                "pmid": item.get("pmid"),
                "conclusion": item.get("core_conclusion"),
                "deleted": bool(item.get("_deleted")),
            }
            for item in items
        ],
        "attacks": [(a["row"], a["dimension"], a["issue"], a["severity"]) for a in attacks],
    }
    return json.dumps(payload, sort_keys=True, ensure_ascii=False)


def _write_table(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_cell_value(row.get(header, "")) for header in headers])


def _cell_value(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)
