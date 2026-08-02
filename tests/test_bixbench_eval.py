import json
import zipfile
from pathlib import Path

from pico.providers.clients import FakeModelClient

from corecoder.eval.bixbench import _newick_tree_metrics, _score_official_answer, run_bixbench, run_official_bixbench


def _official_row(question_id, question, ideal, *, eval_mode="str_verifier"):
    return {
        "id": question_id,
        "question": question,
        "ideal": ideal,
        "distractors": ["wrong"],
        "answer": True,
        "categories": "RNA-seq",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": eval_mode,
        "question_id": question_id,
        "short_id": "bix-demo",
        "result": f"The answer is {ideal}",
    }


def test_bixbench_smoke_runner_writes_reports(tmp_path):
    report = run_bixbench(output_dir=tmp_path)

    assert report["runner"] == "bixbench-compatible-v1"
    assert report["benchmark"]["official_bixbench_dataset"] is False
    assert report["summary"]["total_tasks"] == 2
    assert report["summary"]["failed"] == 0
    assert Path(report["artifacts"]["report_json"]).is_file()
    assert Path(report["artifacts"]["report_md"]).is_file()

    persisted = json.loads(Path(report["artifacts"]["report_json"]).read_text(encoding="utf-8"))
    assert persisted["summary"]["pass_rate"] == 1.0


def test_bixbench_string_verifier_handles_thousands_separator():
    score = _score_official_answer("4510", "4,550", "str_verifier")

    assert score["passed"] is True
    assert score["numeric_match"] is True


def test_bixbench_llm_verifier_handles_approximate_fold_language():
    score = _score_official_answer("2.18", "2x larger in fungi", "llm_verifier")

    assert score["passed"] is True
    assert score["numeric_match"] is True


def test_bixbench_runner_rejects_unknown_tool(tmp_path):
    bad = tmp_path / "bad.json"
    bad.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "tasks": [
                    {
                        "id": "bad_tool",
                        "prompt": "bad",
                        "tool": "missing_tool",
                        "args": {},
                        "expect": [{"path": "", "op": "exists"}],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    try:
        run_bixbench(benchmark_path=bad, output_dir=tmp_path / "out")
        raise AssertionError("unknown tool was accepted")
    except ValueError as exc:
        assert "unknown tool" in str(exc)


def test_phylo_newick_metrics_for_verified50_parser():
    metrics = _newick_tree_metrics("((A:1,B:1):0.5,C:3,D:1);")

    assert metrics is not None
    assert round(metrics["tree_length"], 4) == 6.5
    assert round(metrics["evo_rate"], 4) == 1.625
    assert round(metrics["mean_patristic_distance"], 4) == 3.3333
    assert round(metrics["median_patristic_distance"], 4) == 3.25
    assert round(sorted(metrics["long_branch_scores"])[0], 4) == -10.0


def test_official_bixbench_agent_extracts_gene_specific_median_patristic_distance(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    inner_zip = tmp_path / "scogs_fungi.zip"
    animals_zip = tmp_path / "scogs_animals.zip"
    newick = "(Fungi_Ncra:1.31614302235,Fungi_Scer:1.31614302235,Fungi_Spom:1.31614302235,Fungi_Umay:1.31614302235);"
    with zipfile.ZipFile(inner_zip, "w") as zf:
        zf.writestr("981902at2759.faa.mafft.clipkit.treefile", newick)
    with zipfile.ZipFile(animals_zip, "w") as zf:
        zf.writestr("981902at2759.faa.mafft.clipkit.treefile", "(A:1,B:1,C:1,D:1);")
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.write(inner_zip, "data/scogs_fungi.zip")
        zf.write(animals_zip, "data/scogs_animals.zip")
    row = {
        "id": "uuid",
        "question": "What is the median patristic distance for the fungal gene 981902at2759?",
        "ideal": "2.63",
        "distractors": ["1.24"],
        "answer": True,
        "categories": "Phylogenetics and Evolutionary Analysis",
        "paper": "Not Available",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "llm_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 2.63",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    def factory():
        calls.append("called")
        return FakeModelClient(["<final>ANSWER: wrong</final>"])

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        model_client_factory=factory,
        extract_capsules=True,
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "2.6323"


def test_official_bixbench_agent_extracts_gene_specific_evolutionary_rate(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    animals_zip = tmp_path / "scogs_animals.zip"
    fungi_zip = tmp_path / "scogs_fungi.zip"
    with zipfile.ZipFile(animals_zip, "w") as zf:
        zf.writestr("156083at2759.faa.mafft.clipkit.treefile", "(A:0.05,B:0.05,C:0.05,D:0.0384);")
    with zipfile.ZipFile(fungi_zip, "w") as zf:
        zf.writestr("156083at2759.faa.mafft.clipkit.treefile", "(F1:1,F2:1,F3:1,F4:1);")
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.write(animals_zip, "data/scogs_animals.zip")
        zf.write(fungi_zip, "data/scogs_fungi.zip")
    row = {
        "id": "uuid",
        "question": (
            "Calculate the evolutionary rate for the BUSCO gene 156083at2759 "
            "using PhyKIT's evoluionary_rate function. What is the gene's evolutionary rate in animals?"
        ),
        "ideal": "0.0471",
        "distractors": ["0.5"],
        "answer": True,
        "categories": "Phylogenetics and Evolutionary Analysis",
        "paper": "Not Available",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "llm_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 0.0471",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    def factory():
        calls.append("called")
        return FakeModelClient(["<final>ANSWER: wrong</final>"])

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        model_client_factory=factory,
        extract_capsules=True,
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "0.0471"


def test_official_bixbench_agent_extracts_swarm_ratio_closest_to_strain_one(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    csv_text = "\n".join(
        [
            "DishNumber,StrainNumber,Ratio,Replicate,Area,Circularity,Round",
            "1,1,1:0,A,112000,0.061,0.9",
            "2,1,1:0,B,111500,0.062,0.9",
            "3,1,1:0,C,112400,0.061,0.9",
            "4,287_98,2:1,A,87000,0.096,0.9",
            "5,287_98,2:1,B,87100,0.096,0.9",
            "6,287_98,2:1,C,87070,0.097,0.9",
            "7,287_98,5:1,A,118000,0.066,0.9",
            "8,287_98,5:1,B,118100,0.067,0.9",
            "9,287_98,5:1,C,118160,0.067,0.9",
        ]
    )
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/Swarm_2.csv", csv_text)
    row = {
        "id": "uuid",
        "question": (
            "In a bacterial swarming assay compared mixed cultures of Strain 287 and Strain 98 "
            "at various ratios, which ratio produced colonies with a mean area and circularity value "
            "that is most similar to the mean area and circularity values observed in Strain 1?"
        ),
        "ideal": "5:1",
        "distractors": ["2:1"],
        "answer": True,
        "categories": "Imaging",
        "paper": "https://example.test/paper",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 5:1",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    def factory():
        calls.append("called")
        return FakeModelClient(["<final>ANSWER: wrong</final>"])

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        model_client_factory=factory,
        extract_capsules=True,
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "5:1"


def test_official_bixbench_agent_extracts_glm_aic_and_coefficient_from_notebook(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook_text = (
        "Call: glm(formula = Response ~ BMI, family = binomial, data = dataset) "
        "Coefficients: (Intercept) -3.63193 2.15933 -1.682 0.0926 BMI 0.16775 0.09787 1.714 0.0865 "
        "Null deviance: 110.85 on 79 degrees of freedom Residual deviance: 107.81 on 78 degrees of freedom AIC: 111.81 "
        "Call: glm(formula = Response ~ Age, family = binomial, data = dataset) "
        "Coefficients: (Intercept) 4.44657 1.45133 3.064 0.00219 Age -0.07495 0.02435 -3.078 0.00208 "
        "AIC: 104.14"
    )
    notebook = {"cells": [{"cell_type": "markdown", "source": [notebook_text], "outputs": []}]}
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    rows = [
        {
            "id": "uuid-aic",
            "question": "What is the Akaike Information Criterion (AIC) for the logistic regression model that uses BMI as the sole predictor of treatment response (efficacy PR)?",
            "ideal": "(111.80,111.82)",
            "distractors": ["110.85"],
            "answer": True,
            "categories": "Other",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q2",
            "short_id": "bix-demo",
            "result": "The answer is 111.81",
        },
        {
            "id": "uuid-age",
            "question": "In a simple logistic regression model assessing age and camrelizumab treatment response, what is the coefficient estimate (change in log-odds) for age?",
            "ideal": "(-0.064,-0.084)",
            "distractors": ["0.02435"],
            "answer": True,
            "categories": "Other",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q8",
            "short_id": "bix-demo",
            "result": "The answer is -0.07495",
        },
    ]
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    def factory():
        calls.append("called")
        return FakeModelClient(["<final>ANSWER: wrong</final>"])

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        model_client_factory=factory,
        extract_capsules=True,
    )

    assert calls == []
    assert report["summary"]["passed"] == 2
    assert [task["prediction"] for task in report["tasks"]] == ["111.81", "-0.07495"]


def test_official_bixbench_agent_calculates_methylation_filter_statistics(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    jackdaw = "\n".join(
        [
            "Pos,Chromosome,MethylationPercentage",
            "1_10,1,95",
            "1_10,1,97",
            "1_20,1,5",
            "W_10,W,99",
            "W_20,W,50",
        ]
    )
    zebra_finch = "\n".join(
        [
            "Pos,Chromosome,MethylationPercentage",
            "1_10,1,95",
            "1_20,1,50",
            "2_10,2,5",
            "2_20,2,10",
        ]
    )
    lengths = "Chromosome,Length\n1,100\n2,200\nW,25\n"
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/JD_AgeRelated_CpG_noMT_Final.csv", jackdaw)
        zf.writestr("data/JD_Chromosome_Length.csv", lengths)
        zf.writestr("data/ZF_AgeRelated_CpG_noMT_Final.csv", zebra_finch)
        zf.writestr("data/ZF_Chromosome_Length.csv", lengths)
    questions = [
        (
            "What is the mean of per-chromosome densities (CpGs per base pair) for filtered (>90% or <10% methylation) unique age-related CpGs across all chromosomes in the Jackdaw genome?",
            "0.03",
            "range_verifier",
        ),
        ("Which chromosome in the Jackdaw genome shows the highest density of age-related CpG sites?", "Chromosome W", "str_verifier"),
        ("How many individual methylation measurements (rows) are removed when filtering out measurements that do not show >90% or <10% methylation in the Zebra Finch dataset?", "2", "str_verifier"),
    ]
    rows = []
    for index, (question, ideal, eval_mode) in enumerate(questions, start=1):
        rows.append(
            {
                "id": f"uuid-{index}",
                "question": question,
                "ideal": ideal,
                "distractors": ["wrong"],
                "answer": True,
                "categories": "Epigenomics",
                "paper": "",
                "data_folder": "CapsuleFolder-demo.zip",
                "eval_mode": eval_mode,
                "question_id": f"bix-demo-q{index}",
                "short_id": "bix-demo",
                "result": f"The answer is {ideal}",
            }
        )
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 3
    assert [task["prediction"] for task in report["tasks"]] == ["0.03", "Chromosome W", "2"]
    for task in report["tasks"]:
        trace = json.loads(Path(task["artifacts"]["agent_trace"]).read_text(encoding="utf-8"))
        assert trace["method"] == "methylation_density"


def test_official_bixbench_agent_calculates_mageck_replicate_spearman(tmp_path):
    import io
    import openpyxl

    dataset = tmp_path / "official"
    dataset.mkdir()
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "MAGeCK P-values"
    worksheet.append(["RefSeq ID", "Chronic Round1 S1", "Chronic Round1 S2"])
    worksheet.append(["gene-a", 1.0, 1.0])
    worksheet.append(["gene-b", 2.0, 3.0])
    worksheet.append(["gene-c", 3.0, 2.0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr(
            "data/JuliaJong_CRISPRa_BCL2_B3GNT_cacerResTcellCytotoxicity_supplData1_mageck.xlsx",
            buffer.getvalue(),
        )
    row = {
        "id": "uuid",
        "question": "What is the Spearman correlation coefficient between the replicate MAGeCK P-values for chronic round 1?",
        "ideal": "0.5",
        "distractors": ["1.0"],
        "answer": True,
        "categories": "Functional Genomics",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 0.5",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "0.5"
    trace = json.loads(Path(report["tasks"][0]["artifacts"]["agent_trace"]).read_text(encoding="utf-8"))
    assert trace["method"] == "mageck_replicate_correlation"


def test_official_bixbench_agent_extracts_enrichment_metrics_from_notebook_table(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    table = """
    <table>
      <thead><tr><th></th><th>Gene_set</th><th>Term</th><th>Overlap</th><th>Odds Ratio</th></tr></thead>
      <tbody><tr><th>4</th><td>Reactome_2022</td><td>TP53 Regulates Transcription Of Cell Cycle Genes R-HSA-1</td><td>8/49</td><td>6.023533</td></tr></tbody>
    </table>
    """
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": ["enrich_results"],
                "outputs": [{"output_type": "display_data", "data": {"text/html": [table]}}],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    questions = [
        ("What is the odds ratio for p53-mediated cell cycle gene regulation in an enrichment analysis?", "6.023533"),
        ("After enrichment analysis, what is the overlap ratio for the 'TP53 Regulates Transcription Of Cell Cycle Genes' pathway?", "8/49"),
    ]
    rows = []
    for index, (question, ideal) in enumerate(questions, start=1):
        rows.append(
            {
                "id": f"uuid-{index}",
                "question": question,
                "ideal": ideal,
                "distractors": ["wrong"],
                "answer": True,
                "categories": "RNA-seq",
                "paper": "",
                "data_folder": "CapsuleFolder-demo.zip",
                "eval_mode": "str_verifier",
                "question_id": f"bix-demo-q{index}",
                "short_id": "bix-demo",
                "result": f"The answer is {ideal}",
            }
        )
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 2
    assert [task["prediction"] for task in report["tasks"]] == ["6.023533", "8/49"]


def test_official_bixbench_agent_extracts_total_deg_count_from_notebook_output(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": ["# Print out number of significant DEGs\n", "nrow(final_res)\n"],
                "outputs": [
                    {"output_type": "display_data", "data": {"text/plain": ["[1] 2118"]}},
                    {"output_type": "stream", "text": ["[1] 44.95%"]},
                ],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    row = {
        "id": "uuid",
        "question": "What is the total number of significantly differentially expressed genes (padj < 0.05)?",
        "ideal": "2118",
        "distractors": ["952"],
        "answer": True,
        "categories": "RNA-seq",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 2118",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "2118"


def test_official_bixbench_agent_extracts_gene_length_correlation_table(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    table = """
    <table><thead><tr><th></th><th>Pearson correlation</th><th>p_value</th></tr></thead><tbody>
    <tr><th>CD4</th><td>0.049889</td><td>2e-11</td></tr>
    <tr><th>CD8</th><td>0.043670</td><td>5e-09</td></tr>
    <tr><th>CD14</th><td>0.020527</td><td>0.006</td></tr>
    <tr><th>CD19</th><td>0.035662</td><td>1e-06</td></tr>
    </tbody></table>
    """
    notebook = {"cells": [{"cell_type": "code", "source": [], "outputs": [{"data": {"text/html": table}}]}]}
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    questions = [
        ("After calculating Pearson correlations between gene length and average gene expression, in which immune cell type is the correlation weakest?", "CD14"),
        ("What is the Pearson correlation coefficient between gene length and mean expression in CD14 immune cells?", "0.020527"),
    ]
    rows = [_official_row(f"bix-demo-q{i}", question, ideal) for i, (question, ideal) in enumerate(questions, start=1)]
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []
    report = run_official_bixbench(dataset, output_dir=tmp_path / "out", strategy="biocoreagent_agent", extract_capsules=True, model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]))
    assert calls == []
    assert report["summary"]["passed"] == 2
    assert [task["prediction"] for task in report["tasks"]] == ["CD14", "0.020527"]


def test_official_bixbench_agent_extracts_top_pathway_fraction_and_replicate_direction(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    body = "".join(
        f"<tr><th>{i}</th><td>WikiPathways_2019_Mouse</td><td>{'Oxidative damage' if i == 1 else 'Oxidative stress' if i == 5 else 'Pathway '+str(i)}</td></tr>"
        for i in range(20)
    )
    table = f"<table><thead><tr><th></th><th>Gene_set</th><th>Term</th></tr></thead><tbody>{body}</tbody></table>"
    text = "View of AnnData object with n_obs x n_vars = 6 x 1416 View of AnnData object with n_obs x n_vars = 4 x 1942"
    notebook = {"cells": [{"cell_type": "code", "source": [text], "outputs": [{"data": {"text/html": table}}]}]}
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    rows = [
        _official_row("bix-demo-q1", "Using WikiPathways_2019_Mouse, what is the fraction of pathways whose name contains \"oxidative\" among the top 20 enriched pathways?", "0.1"),
        _official_row("bix-demo-q2", "Repeat differential expression analysis excluding the third replicates and describe how this affects the number of significantly differentially expressed genes.", "Increases the number of differentially expressed genes", eval_mode="llm_verifier"),
    ]
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []
    report = run_official_bixbench(dataset, output_dir=tmp_path / "out", strategy="biocoreagent_agent", extract_capsules=True, model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]))
    assert calls == []
    assert [task["prediction"] for task in report["tasks"]] == ["0.1", "Increases the number of differentially expressed genes"]
    assert report["summary"]["passed"] == 2


def test_official_bixbench_agent_extracts_directional_conclusion(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {"cells": [{"cell_type": "markdown", "source": ["The metabolic response is mostly driven by the DOWN-regulated genes."], "outputs": []}]}
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    row = _official_row("bix-demo-q1", "Determine whether upregulation or downregulation of genes primarily drives the metabolic effects.", "Downregulation", eval_mode="llm_verifier")
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []
    report = run_official_bixbench(dataset, output_dir=tmp_path / "out", strategy="biocoreagent_agent", extract_capsules=True, model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]))
    assert calls == []
    assert report["summary"]["passed"] == 1


def test_official_bixbench_agent_extracts_proteomics_fold_change(tmp_path):
    import io
    import openpyxl

    dataset = tmp_path / "official"
    dataset.mkdir()
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Tumor vs Normal"
    worksheet.append(["gene", "Normal", "Tumor", "FC", "log2FC", "compare"])
    worksheet.append(["ENO1", 10.0, 48.1, 4.81, 2.27, "Tumor vs Normal"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/Proteomic_data.xlsx", buffer.getvalue())
    rows = [
        _official_row(
            "bix-demo-q1",
            "Based on the proteomics data, what is the fold change in ENO1 protein abundance between tumor and normal samples?",
            "4.81-fold increase in tumor",
            eval_mode="llm_verifier",
        ),
        _official_row(
            "bix-demo-q2",
            "What is the log2 fold change value rounded to 2 decimal places for ENO1 in the proteomics dataset?",
            "2.27",
        ),
    ]
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []
    report = run_official_bixbench(dataset, output_dir=tmp_path / "out", strategy="biocoreagent_agent", extract_capsules=True, model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]))
    assert calls == []
    assert report["summary"]["passed"] == 2
    assert [task["prediction"] for task in report["tasks"]] == ["4.81-fold increase in tumor", "2.27"]


def test_official_bixbench_agent_counts_oldest_carrier_variants_and_coding_synonymous_fraction(tmp_path):
    import io
    import openpyxl

    dataset = tmp_path / "official"
    dataset.mkdir()
    status_workbook = openpyxl.Workbook()
    status_sheet = status_workbook.active
    status_sheet.append(["Sample ID", "Status", "Sex", "BLM Mutation Status", "Age", "Cancer"])
    status_sheet.append([100, "Father", "M", "Carrier", 40, "N"])
    status_sheet.append([556, "Father", "M", "Carrier", 65, "N"])
    status_buffer = io.BytesIO()
    status_workbook.save(status_buffer)
    status_workbook.close()

    variant_workbook = openpyxl.Workbook()
    variant_sheet = variant_workbook.active
    variant_sheet.append(["Variant Info", None, None, None, None])
    variant_sheet.append(["Zygosity", "Gene Names", "In_CHIP", "Variant Allele Freq", "Sequence Ontology (Combined)"])
    variant_sheet.append(["Reference", "NOTCH1", True, None, "synonymous_variant"])
    variant_sheet.append(["Heterozygous", "NOTCH1", True, 0.1, "synonymous_variant"])
    variant_sheet.append(["Heterozygous", "NOTCH1", True, 0.2, "synonymous_variant"])
    variant_sheet.append(["Heterozygous", "ASXL1", True, 0.2, "missense_variant"])
    variant_sheet.append(["Heterozygous", "SF3B1", True, 0.2, "splice_region_variant"])
    variant_buffer = io.BytesIO()
    variant_workbook.save(variant_buffer)
    variant_workbook.close()

    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/230215_Trio_Status.xlsx", status_buffer.getvalue())
        zf.writestr("data/CHIP_DP10_GQ20_PASS/230209_Exome_GRCh38_CHIP_556-PF.xlsx", variant_buffer.getvalue())
    rows = [
        _official_row("bix-demo-q1", "Which gene has the most non-reference variants in the oldest male carrier?", "NOTCH1"),
        _official_row("bix-demo-q2", "In the BLM mutation carrier cohort, what fraction of coding variants with a VAF below 0.3 are annotated as synonymous?", "(0.65,0.68)", eval_mode="range_verifier"),
    ]
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []
    report = run_official_bixbench(dataset, output_dir=tmp_path / "out", strategy="biocoreagent_agent", extract_capsules=True, model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]))
    assert calls == []
    assert report["summary"]["passed"] == 2
    assert report["tasks"][0]["prediction"] == "NOTCH1"


def test_official_bixbench_blind_baseline_scores_without_leaking_answers(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/table.csv", "x,y\n1,2\n")
    row = {
        "id": "uuid",
        "question": "What p-value did the analysis produce?",
        "ideal": "0.0002",
        "distractors": ["0.5"],
        "answer": True,
        "categories": "RNA-seq",
        "paper": "https://example.test/paper",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 0.0002",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")

    report = run_official_bixbench(dataset, output_dir=tmp_path / "out")

    assert report["benchmark"]["official_bixbench_dataset"] is True
    assert report["summary"]["total_tasks"] == 1
    assert report["summary"]["pass_at_1"] == 0.0
    task = report["tasks"][0]
    assert task["prediction"] == "INSUFFICIENT_EVIDENCE"
    manifest = Path(task["workspace"]) / "blind_baseline_manifest.json"
    assert manifest.exists()
    manifest_text = manifest.read_text(encoding="utf-8")
    assert "0.0002" not in manifest_text
    assert "The answer is 0.0002" not in manifest_text


def test_official_bixbench_auto_detects_verified_50_jsonl(tmp_path):
    dataset = tmp_path / "verified"
    dataset.mkdir()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/table.csv", "x,y\n1,2\n")
    row = {
        "id": "uuid",
        "question": "What p-value did the analysis produce?",
        "ideal": "0.0002",
        "distractors": ["0.5"],
        "answer": True,
        "categories": "RNA-seq",
        "paper": "https://example.test/paper",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 0.0002",
    }
    (dataset / "BixBench-Verified-50.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")

    report = run_official_bixbench(dataset, output_dir=tmp_path / "out")

    assert report["benchmark"]["name"] == "phylobio/BixBench-Verified-50"
    assert report["benchmark"]["verified_50"]["computed"] is True
    assert report["benchmark"]["task_count"] == 1


def test_official_bixbench_filters_by_short_id(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/table.csv", "x,y\n1,2\n")
    rows = []
    for index, short_id in enumerate(["bix-a", "bix-b"], start=1):
        rows.append(
            {
                "id": f"uuid-{index}",
                "question": "What p-value did the analysis produce?",
                "ideal": "0.0002",
                "distractors": ["0.5"],
                "answer": True,
                "categories": "RNA-seq",
                "paper": "https://example.test/paper",
                "data_folder": "CapsuleFolder-demo.zip",
                "eval_mode": "str_verifier",
                "question_id": f"{short_id}-q1",
                "short_id": short_id,
                "result": "The answer is 0.0002",
            }
        )
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")

    report = run_official_bixbench(dataset, output_dir=tmp_path / "out", include_short_ids=["bix-b"])

    assert report["summary"]["total_tasks"] == 1
    assert report["tasks"][0]["id"] == "bix-b-q1"


def test_official_bixbench_agent_strategy_scores_and_writes_artifacts(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/table.csv", "x,y\n1,2\n")
    row = {
        "id": "uuid",
        "question": "What p-value did the analysis produce?",
        "ideal": "0.0002",
        "distractors": ["0.5"],
        "answer": True,
        "categories": "RNA-seq",
        "paper": "https://example.test/paper",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 0.0002",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    clients = []

    def factory():
        client = FakeModelClient(["<final>I inspected the capsule.\nANSWER: 0.0002</final>"])
        clients.append(client)
        return client

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        model_client_factory=factory,
        task_timeout_seconds=30,
    )

    assert report["summary"]["passed"] == 1
    task = report["tasks"][0]
    assert task["passed"] is True
    assert Path(task["artifacts"]["agent_answer"]).exists()
    assert Path(task["artifacts"]["agent_trace"]).exists()
    assert Path(task["artifacts"]["score"]).exists()
    assert Path(task["artifacts"]["task_report"]).exists()
    prompt = clients[0].prompts[0]
    assert "What p-value did the analysis produce?" in prompt
    assert "0.0002" not in prompt
    assert "The answer is 0.0002" not in prompt


def test_official_bixbench_agent_extracts_executed_notebook_statistics(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "markdown",
                "source": [
                    "Looking at our results above we can see that ",
                    "in expect_interact_cat the odds_ratio = 0.754981 , p-val = 5.163146e-02.",
                ],
                "outputs": [],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    row = {
        "id": "uuid",
        "question": (
            "Using an ordinal logistic regression model, what is the odds ratio associated "
            "with healthcare workers' expected patient interaction (expect_interact_cat)?"
        ),
        "ideal": "(0.74,0.77)",
        "distractors": ["1.52"],
        "answer": True,
        "categories": "Other",
        "paper": "https://example.test/paper",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "range_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 0.754981",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    def factory():
        calls.append("called")
        return FakeModelClient(["<final>ANSWER: wrong</final>"])

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=factory,
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    task = report["tasks"][0]
    assert task["prediction"] == "0.754981"
    trace = json.loads(Path(task["artifacts"]["agent_trace"]).read_text(encoding="utf-8"))
    assert trace["strategy"] == "deterministic_capsule_parser"
    assert trace["method"] == "notebook_stats"


def test_official_bixbench_agent_calculates_reduction_from_executed_notebook_or(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "markdown",
                "source": ["expect_interact_cat unexpectedly had odds_ratio = 0.754981 , p-val = 5.163146e-02."],
                "outputs": [],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    row = {
        "id": "uuid",
        "question": (
            "What is the percentage reduction in odds ratio for higher severity among healthcare "
            "workers expected to interact with patients versus those who do not?"
        ),
        "ideal": "(24,26)",
        "distractors": ["1.52"],
        "answer": True,
        "categories": "Other",
        "paper": "https://example.test/paper",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "range_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "The answer is 24.5",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "24.50"


def test_official_bixbench_agent_extracts_scogs_treeness(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()

    def write_scogs_zip(zf, archive_name, values):
        with zipfile.ZipFile(tmp_path / archive_name, "w") as inner:
            for index, value in enumerate(values, start=1):
                internal = value
                terminal = (1.0 - value) / 2.0
                inner.writestr(
                    f"gene{index}.faa.mafft.clipkit.treefile",
                    f"(A:{terminal},B:{terminal}):{internal};",
                )
                inner.writestr(
                    f"gene{index}.faa.mafft.clipkit.iqtree",
                    "Input data: 2 sequences with 100 amino-acid sites\n"
                    f"Number of parsimony informative sites: {index}\n",
                )
        zf.write(tmp_path / archive_name, f"data/{archive_name}")

    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        write_scogs_zip(zf, "scogs_animals.zip", [0.0, 0.0, 0.02])
        write_scogs_zip(zf, "scogs_fungi.zip", [0.04, 0.06, 0.08])

    row = {
        "id": "uuid",
        "question": "What is the median treeness value for fungal genes?",
        "ideal": "0.0600",
        "distractors": [],
        "answer": True,
        "categories": "Genomics,Phylogenetics and Evolutionary Analysis",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "0.0600",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "0.0600"


def test_official_bixbench_agent_extracts_scogs_parsimony_from_alignments(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()

    def alignment_with_informative_sites(count):
        columns = ["AABB" for _ in range(count)] + ["AAAA" for _ in range(10 - count)]
        seqs = ["".join(column[i] for column in columns) for i in range(4)]
        return "\n".join(f">s{i}\n{seq}" for i, seq in enumerate(seqs, start=1))

    def write_scogs_zip(zf, archive_name, counts):
        with zipfile.ZipFile(tmp_path / archive_name, "w") as inner:
            for index, count in enumerate(counts, start=1):
                inner.writestr(f"gene{index}.faa.mafft", alignment_with_informative_sites(count))
        zf.write(tmp_path / archive_name, f"data/{archive_name}")

    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        write_scogs_zip(zf, "scogs_animals.zip", [0, 1, 2])
        write_scogs_zip(zf, "scogs_fungi.zip", [1, 3, 5])

    row = {
        "id": "uuid",
        "question": "What is the median percentage of parsimony informative sites across fungal gene alignments?",
        "ideal": "30.0%",
        "distractors": [],
        "answer": True,
        "categories": "Genomics,Phylogenetics and Evolutionary Analysis",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "30.0%",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "30.0%"


def test_official_bixbench_agent_extracts_scogs_saturation_from_alignments(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()

    def alignment_with_pairwise_distance(distance):
        length = 100
        changed = int(distance * length)
        reference = "A" * length
        variant = ("B" * changed) + ("A" * (length - changed))
        return f">s1\n{reference}\n>s2\n{variant}\n"

    def write_scogs_zip(zf, archive_name, distances):
        with zipfile.ZipFile(tmp_path / archive_name, "w") as inner:
            for index, distance in enumerate(distances, start=1):
                inner.writestr(f"gene{index}.faa.mafft", alignment_with_pairwise_distance(distance))
        zf.write(tmp_path / archive_name, f"data/{archive_name}")

    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        write_scogs_zip(zf, "scogs_animals.zip", [0.10, 0.20, 0.30])
        write_scogs_zip(zf, "scogs_fungi.zip", [0.50, 0.62, 0.80])

    row = {
        "id": "uuid",
        "question": "What is the median saturation value for fungal genes?",
        "ideal": "0.62",
        "distractors": [],
        "answer": True,
        "categories": "Whole Genome Sequencing (WGS),Phylogenetics and Evolutionary Analysis",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "0.62",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "0.62"


def test_official_bixbench_agent_extracts_crispr_correlation_from_notebook(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": [],
                "outputs": [
                    {
                        "name": "stdout",
                        "output_type": "stream",
                        "text": [
                            "Bottom 5 negatively correlated genes:\n",
                            "                  gene  spearman_r       p_value   adj_p_value  significant\n",
                            "2537     CDKN1A (1026)   -0.523398  1.332245e-78  4.709752e-75         True\n",
                            "Proportion of genes with significant positive correlation: 14.25%\n",
                            "Proportion of genes with significant negative correlation: 7.56%\n",
                            "Only three genes show strong positive correlation with a Spearman coefficient >= 0.6.\n",
                            "Both datasets appear to be skewed and considerably deviate from normal distribution.\n",
                        ],
                    }
                ],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/CRISPRGeneEffect.csv", ",A\nACH-1,1\n")
        zf.writestr("data/OmicsExpressionProteinCodingGenesTPMLogp1BatchCorrected.csv", ",A\nACH-1,1\n")
        zf.writestr("data/Model.csv", "ModelID\nACH-1\n")
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    row = {
        "id": "uuid",
        "question": "In the provided data, what gene symbol has the strongest negative Spearman correlation between its expression and essentiality?",
        "ideal": "CDKN1A",
        "distractors": [],
        "answer": True,
        "categories": "Transcriptomics,Functional Genomics",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "CDKN1A",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "CDKN1A"
    trace = json.loads(Path(report["tasks"][0]["artifacts"]["agent_trace"]).read_text(encoding="utf-8"))
    assert trace["method"] == "crispr_correlation"


def test_official_bixbench_agent_formats_crispr_correlation_range_answers(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": [],
                "outputs": [
                    {
                        "name": "stdout",
                        "output_type": "stream",
                        "text": [
                            "Proportion of genes with significant positive correlation: 14.25%\n",
                            "Proportion of genes with significant negative correlation: 7.56%\n",
                            "Both datasets appear to be skewed and considerably deviate from normal distribution.\n",
                        ],
                    }
                ],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/CRISPRGeneEffect.csv", ",A\nACH-1,1\n")
        zf.writestr("data/OmicsExpressionProteinCodingGenesTPMLogp1BatchCorrected.csv", ",A\nACH-1,1\n")
        zf.writestr("data/Model.csv", "ModelID\nACH-1\n")
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    rows = [
        {
            "id": "uuid-1",
            "question": "Using a Spearman Rank correlation test and Benjamini-Hochberg multi-test correction, what percentage of genes show a statistically significant correlation between expression and essentiality, in either direction?",
            "ideal": "(20,25)",
            "distractors": [],
            "answer": True,
            "categories": "Transcriptomics,Functional Genomics",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q1",
            "short_id": "bix-demo",
            "result": "21.81",
        },
        {
            "id": "uuid-2",
            "question": "What is the skewness characteristic of the gene expression (log2(TPM+1)) distribution across all cell lines in the dataset?",
            "ideal": "Right-skewed with a long tail",
            "distractors": [],
            "answer": True,
            "categories": "Transcriptomics,Functional Genomics",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "eval_mode": "str_verifier",
            "question_id": "bix-demo-q2",
            "short_id": "bix-demo",
            "result": "Right-skewed with a long tail",
        },
    ]
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 2
    assert [task["prediction"] for task in report["tasks"]] == ["21.81", "Right-skewed with a long tail"]


def test_official_bixbench_agent_extracts_variant_pathogenicity_table(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": [],
                "outputs": [
                    {
                        "name": "stdout",
                        "output_type": "stream",
                        "text": [
                            "Benign 0.92857143 BSyn Probands\n",
                            "Likely Benign 0.07142857 BSyn Probands\n",
                            "Benign 0.94594595 BLM Carriers\n",
                            "Likely Benign 0.05405405 BLM Carriers\n",
                            "Benign 1.00000000 Control Children\n",
                            "Likely Benign 0.00000000 Control Children\n",
                            "Benign 1.00000000 Control Parents\n",
                            "Likely Benign 0.00000000 Control Parents\n",
                        ],
                    }
                ],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/230215_Trio_Status.xlsx", "placeholder")
        zf.writestr("data/230214_Schenz_et_al_2022_CHIP_Genes.xlsx", "placeholder")
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    row = {
        "id": "uuid",
        "question": "Among samples with a 'Carrier' BLM mutation status, what proportion of somatic CHIP variants (VAF < 0.3) can be classified as benign?",
        "ideal": "(0.90, 1.00)",
        "distractors": [],
        "answer": True,
        "categories": "Genomics,Genomic Variant Analysis",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "range_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "0.94594595",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 1
    assert report["tasks"][0]["prediction"] == "0.94594595"


def test_official_bixbench_agent_extracts_swarm_imaging_stats(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    swarm_csv = "\n".join([
        "DishNumber,Genotype,StrainNumber,Ratio,Replicate,Area,Circularity,Round",
        "1,ΔlasR,300,1:0,1,19193,0.409,0.978",
        "2,ΔlasR,300,1:0,2,20556,0.448,0.639",
        "3,ΔlasR,300,1:0,3,32011,0.403,0.872",
        "7,rhlR-,61,1:0,1,10516,0.509,0.978",
        "8,rhlR-,61,1:0,2,11220,0.601,0.974",
        "9,rhlR-,61,1:0,3,10122,0.606,0.973",
        "13,Wildtype,1,1:0,1,95704,0.076,0.839",
        "14,Wildtype,1,1:0,2,69298,0.09,0.87",
        "15,Wildtype,1,1:0,3,82325,0.062,0.948",
    ])
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/Swarm_1.csv", swarm_csv)
    rows = [
        {
            "id": "uuid-1",
            "question": "What is the approximate percent reduction in mean colony area for the ΔlasR mutant compared to wildtype?",
            "ideal": "(69,72)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q1",
            "short_id": "bix-demo",
        },
        {
            "id": "uuid-2",
            "question": "What is the relative proportion of mean colony area of the ΔlasR mutant to the wildtype strain?",
            "ideal": "(25,30)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q2",
            "short_id": "bix-demo",
        },
    ]
    for row in rows:
        row.update({
            "distractors": [],
            "answer": True,
            "categories": "Imaging",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "result": "",
        })
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 2
    assert [task["prediction"] for task in report["tasks"]] == ["70.99", "29.01"]


def test_official_bixbench_agent_extracts_neun_notebook_stats(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": [],
                "outputs": [
                    {
                        "name": "stdout",
                        "output_type": "stream",
                        "text": [
                            "Shapiro-Wilk Test for KD samples: ShapiroResult(statistic=0.9563693455884582, pvalue=0.7748936950981364)\n",
                            "Cohen's d: 0.21628441720382488\n",
                            "Required sample size per group: 337.0\n",
                            "C(Hemisphere):C(Sex) 333.0625 1.0 1.065587 0.322301\n",
                        ],
                    }
                ],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/NeuN_quantification.csv", "Sample,Hemispere,NeuN,Sex\n1,KD,221,M\n1,CTRL,237,M\n")
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    rows = [
        {
            "id": "uuid-1",
            "question": "What is the Cohen's d effect size for the difference in NeuN counts between KD and CTRL conditions in the dataset?",
            "ideal": "(0.215,0.217)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q1",
            "short_id": "bix-demo",
        },
        {
            "id": "uuid-2",
            "question": "Based on a Shapiro-Wilk test, what is the W statistic for evaluating whether the NeuN counts in the KD hemisphere group follow a normal distribution?",
            "ideal": "(0.955,0.957)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q2",
            "short_id": "bix-demo",
        },
        {
            "id": "uuid-3",
            "question": "What minimum sample size per group would be required to detect a statistically significant difference in NeuN counts between KD and CTRL groups, given the observed effect size and desired power of 0.8?",
            "ideal": "337 samples",
            "eval_mode": "llm_verifier",
            "question_id": "bix-demo-q3",
            "short_id": "bix-demo",
        },
    ]
    for row in rows:
        row.update({
            "distractors": [],
            "answer": True,
            "categories": "Imaging,Other",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "result": "",
        })
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 3
    assert [task["prediction"] for task in report["tasks"]] == ["0.216284", "0.956369", "337"]


def test_official_bixbench_agent_extracts_bcg_corona_chisquare(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "source": [],
                "outputs": [
                    {
                        "name": "stdout",
                        "output_type": "stream",
                        "text": [
                            "Chi-square test for expect_interract = Yes\n",
                            "Chi2=9.31508641975309, p=0.025382121082660535, dof=3\n",
                            "Chi-square test for patients_seen = 1-50\n",
                            "Chi2=9.420743463606744, p=0.024189637931453605, dof=3\n",
                            "Chi-square test for patients_seen = >100\n",
                            "Chi2=2.6785714285714284, p=0.4438811470106675, dof=3\n",
                        ],
                    }
                ],
            }
        ]
    }
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/TASK008_BCG-CORONA_AE.csv", "USUBJID,TRTGRP,AESEV\nA,BCG,1\n")
        zf.writestr("data/TASK008_BCG-CORONA_DM.csv", "USUBJID,patients_seen,expect_interact\nA,1-50,Yes\n")
        zf.writestr("data/TASK008_BCG-CORONA_EX.csv", "USUBJID,TRTGRP\nA,BCG\n")
        zf.writestr("notebook/executed.ipynb", json.dumps(notebook))
    rows = [
        {
            "id": "uuid-1",
            "question": "Using a chi-square test, what is the p-value for the association between BCG vaccination and COVID-19 severity among healthcare workers who expect to interact with patients?",
            "ideal": "(0.024,0.026)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q1",
            "short_id": "bix-demo",
        },
        {
            "id": "uuid-2",
            "question": "For healthcare workers seeing over 100 patients, what is the statistical significance (p-value) of the association between vaccination status and disease severity?",
            "ideal": "(0.43,0.45)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q2",
            "short_id": "bix-demo",
        },
        {
            "id": "uuid-3",
            "question": "When performing a Chi-square test of independence to examine the association between BCG vaccination status (BCG vs Placebo) and COVID-19 disease severity among healthcare workers who see 1-50 patients, what is the p-value of the statistical test?",
            "ideal": "(0.023,0.025)",
            "eval_mode": "range_verifier",
            "question_id": "bix-demo-q3",
            "short_id": "bix-demo",
        },
    ]
    for row in rows:
        row.update({
            "distractors": [],
            "answer": True,
            "categories": "Other",
            "paper": "",
            "data_folder": "CapsuleFolder-demo.zip",
            "result": "",
        })
    (dataset / "BixBench.jsonl").write_text("\n".join(json.dumps(row) for row in rows) + "\n", encoding="utf-8")
    calls = []

    report = run_official_bixbench(
        dataset,
        output_dir=tmp_path / "out",
        strategy="biocoreagent_agent",
        extract_capsules=True,
        model_client_factory=lambda: calls.append("called") or FakeModelClient(["<final>ANSWER: wrong</final>"]),
    )

    assert calls == []
    assert report["summary"]["passed"] == 3
    assert [task["prediction"] for task in report["tasks"]] == ["0.025382", "0.443881", "0.024190"]


def test_official_bixbench_resume_reuses_task_report(tmp_path):
    dataset = tmp_path / "official"
    dataset.mkdir()
    with zipfile.ZipFile(dataset / "CapsuleFolder-demo.zip", "w") as zf:
        zf.writestr("data/table.csv", "x,y\n1,2\n")
    row = {
        "id": "uuid",
        "question": "What is the answer?",
        "ideal": "abc",
        "distractors": [],
        "answer": True,
        "categories": "Other",
        "paper": "",
        "data_folder": "CapsuleFolder-demo.zip",
        "eval_mode": "str_verifier",
        "question_id": "bix-demo-q1",
        "short_id": "bix-demo",
        "result": "abc",
    }
    (dataset / "BixBench.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
    out = tmp_path / "out"
    first = run_official_bixbench(
        dataset,
        output_dir=out,
        strategy="biocoreagent_agent",
        model_client_factory=lambda: FakeModelClient(["<final>ANSWER: abc</final>"]),
    )
    second = run_official_bixbench(
        dataset,
        output_dir=out,
        strategy="biocoreagent_agent",
        model_client_factory=lambda: FakeModelClient(["<final>ANSWER: wrong</final>"]),
        resume=True,
    )

    assert first["tasks"][0]["prediction"] == "abc"
    assert second["tasks"][0]["prediction"] == first["tasks"][0]["prediction"]
